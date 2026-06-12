# excel_exporter.py
# Excel export functionality
# Copyright (c) 2026 Tran Xuan An

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from shape_calculator import calculate_perimeter, calculate_statistics


class ExcelExporter:
    """Export polygon data to Excel with formatted sheets"""
    
    def __init__(self, polygons, output_folder):
        self.polygons = polygons
        self.output_folder = output_folder
        
    def export(self):
        """Export data to Excel file"""
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_file = os.path.join(
                self.output_folder,
                f"Column_Inspector_{timestamp}.xlsx"
            )
            
            # Create workbook
            wb = Workbook()
            
            # Create Summary sheet first
            ws_summary = wb.active
            ws_summary.title = "Summary"
            self._create_summary_sheet(ws_summary)
            
            # Create Detail sheet
            ws_detail = wb.create_sheet("Shape Details")
            self._create_detail_sheet(ws_detail)
            
            # Create batch files for shape finding
            self._create_batch_files()
            
            # Save workbook
            wb.save(output_file)
            print(f"\n✓ Excel file created with {len(self.polygons)} shape(s): {output_file}")
            print(f"✓ Summary sheet created with statistics by name")
            print(f"✓ Shape reverse lookup enabled - click '🔍 Find in CAD' buttons in Detail sheet")
            
            return output_file
            
        except Exception as e:
            print(f"✗ Error creating Excel file: {e}")
            return None
    
    def _create_batch_files(self):
        """Create VBScript files for each shape to enable clicking from Excel (hidden console)"""
        try:
            # Create scripts subfolder
            scripts_folder = os.path.join(self.output_folder, "scripts")
            if not os.path.exists(scripts_folder):
                os.makedirs(scripts_folder)
                print(f"✓ Created scripts folder: {scripts_folder}")
            
            # Find FindShape executable or Python script
            finder_exe, is_executable = self._find_shape_finder()
            
            for polygon in self.polygons:
                shape_id = polygon.get('id', '')
                if shape_id:
                    # Create VBScript file in scripts folder
                    vbs_file = os.path.join(scripts_folder, f"find_{shape_id}.vbs")
                    
                    with open(vbs_file, 'w') as f:
                        # VBScript to run without showing console window
                        f.write('Set objShell = CreateObject("WScript.Shell")\n')
                        if is_executable:
                            # Direct executable call
                            f.write(f'objShell.Run "\""{finder_exe}\"" {shape_id}", 0, False\n')
                        else:
                            # Python script call
                            python_exe = self._find_python()
                            f.write(f'objShell.Run "\""{python_exe}\"" \""{finder_exe}\"" {shape_id}", 0, False\n')
                    
                    # Also create a simpler .bat file for manual testing if needed
                    batch_file = os.path.join(scripts_folder, f"find_{shape_id}.bat")
                    with open(batch_file, 'w') as f:
                        f.write('@echo off\n')
                        f.write(f'cd /d "{self.output_folder}"\n')
                        if is_executable:
                            f.write(f'"{finder_exe}" {shape_id}\n')
                        else:
                            f.write(f'python "{finder_exe}" {shape_id}\n')
                        # No pause - auto close after execution
        except Exception as e:
            print(f"⚠ Warning: Could not create batch files: {e}")
    
    def _find_shape_finder(self):
        """Find FindShape.exe or find_shape.py
        Returns: (path, is_executable)
        """
        # Priority 1: FindShape.exe in same folder as output
        exe_path = os.path.abspath(os.path.join(self.output_folder, "FindShape.exe"))
        if os.path.exists(exe_path):
            print(f"✓ Using FindShape.exe (compiled, no Python required)")
            return (exe_path, True)
        
        # Priority 2: FindShape.exe in current directory
        exe_path = os.path.abspath("FindShape.exe")
        if os.path.exists(exe_path):
            print(f"✓ Using FindShape.exe from current directory")
            return (exe_path, True)
        
        # Priority 3: FindShape.exe in PATH
        try:
            result = os.system('where FindShape.exe >nul 2>&1')
            if result == 0:
                print(f"✓ Using FindShape.exe from PATH")
                return ("FindShape.exe", True)
        except:
            pass
        
        # Fallback: find_shape.py
        py_path = os.path.abspath(os.path.join(self.output_folder, "find_shape.py"))
        if os.path.exists(py_path):
            print(f"⚠ Using find_shape.py (requires Python)")
            return (py_path, False)
        
        # Last resort: try current directory
        py_path = os.path.abspath("find_shape.py")
        print(f"⚠ Using find_shape.py from current directory")
        return (py_path, False)
    
    def _find_python(self):
        """Find Python executable path"""
        # Try common locations
        python_paths = [
            'python',
            'python3',
            r'C:\Python39\python.exe',
            r'C:\Python310\python.exe',
            r'C:\Python311\python.exe',
            r'C:\Python312\python.exe',
            r'C:\Program Files\Python39\python.exe',
            r'C:\Program Files\Python310\python.exe',
            r'C:\Program Files\Python311\python.exe',
        ]
        
        # Check if python is in PATH
        for py in python_paths:
            try:
                result = os.system(f'where {py} >nul 2>&1')
                if result == 0:
                    return py
            except:
                pass
        
        return 'python'  # Fallback to default
    
    def _create_summary_sheet(self, ws):
        """Create summary statistics sheet"""
        stats = calculate_statistics(self.polygons)
        
        # Styles
        title_font = Font(bold=True, size=16, color="FFFFFF")
        title_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:G1')
        cell = ws.cell(row=1, column=1)
        cell.value = "SHAPE STATISTICS SUMMARY"
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30
        
        # Headers
        headers = ["Name", "Count", "Total Area", "Avg Area", "Total Perimeter", "Avg Perimeter", "Sample ID"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        
        # Data
        row = 4
        for name in sorted(stats.keys()):
            stat = stats[name]
            avg_area = stat['total_area'] / stat['count']
            avg_perimeter = stat['total_perimeter'] / stat['count']
            
            # Get first shape ID with this name
            sample_id = next((p.get('id', 'N/A') for p in self.polygons if p.get('name') == name), 'N/A')
            
            ws.cell(row=row, column=1).value = name
            ws.cell(row=row, column=1).border = border
            
            ws.cell(row=row, column=2).value = stat['count']
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2).border = border
            
            ws.cell(row=row, column=3).value = round(stat['total_area'], 2)
            ws.cell(row=row, column=3).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=3).border = border
            
            ws.cell(row=row, column=4).value = round(avg_area, 2)
            ws.cell(row=row, column=4).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=4).border = border
            
            ws.cell(row=row, column=5).value = round(stat['total_perimeter'], 2)
            ws.cell(row=row, column=5).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=5).border = border
            
            ws.cell(row=row, column=6).value = round(avg_perimeter, 2)
            ws.cell(row=row, column=6).alignment = Alignment(horizontal="right")
            ws.cell(row=row, column=6).border = border
            
            ws.cell(row=row, column=7).value = sample_id
            ws.cell(row=row, column=7).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=7).border = border
            ws.cell(row=row, column=7).font = Font(color="0000FF", underline="single")
            
            row += 1
        
        # Totals row
        self._add_totals_row(ws, row, stats, border)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 20
    
    def _add_totals_row(self, ws, row, stats, border):
        """Add totals row to summary sheet"""
        row += 1
        ws.cell(row=row, column=1).value = "TOTAL"
        ws.cell(row=row, column=1).font = Font(bold=True)
        ws.cell(row=row, column=1).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=1).border = border
        
        total_count = sum(s['count'] for s in stats.values())
        total_area = sum(s['total_area'] for s in stats.values())
        total_perimeter = sum(s['total_perimeter'] for s in stats.values())
        
        ws.cell(row=row, column=2).value = total_count
        ws.cell(row=row, column=2).font = Font(bold=True)
        ws.cell(row=row, column=2).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2).border = border
        
        ws.cell(row=row, column=3).value = round(total_area, 2)
        ws.cell(row=row, column=3).font = Font(bold=True)
        ws.cell(row=row, column=3).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=3).border = border
        
        ws.cell(row=row, column=4).value = round(total_area / total_count if total_count > 0 else 0, 2)
        ws.cell(row=row, column=4).font = Font(bold=True)
        ws.cell(row=row, column=4).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=4).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=4).border = border
        
        ws.cell(row=row, column=5).value = round(total_perimeter, 2)
        ws.cell(row=row, column=5).font = Font(bold=True)
        ws.cell(row=row, column=5).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=5).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=5).border = border
        
        ws.cell(row=row, column=6).value = round(total_perimeter / total_count if total_count > 0 else 0, 2)
        ws.cell(row=row, column=6).font = Font(bold=True)
        ws.cell(row=row, column=6).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=6).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=6).border = border
        
        ws.cell(row=row, column=7).value = "-"
        ws.cell(row=row, column=7).font = Font(bold=True)
        ws.cell(row=row, column=7).fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        ws.cell(row=row, column=7).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=7).border = border
    
    def _create_detail_sheet(self, ws):
        """Create detailed shape data sheet"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        
        row = 1
        
        # Process each polygon
        for poly_idx, polygon in enumerate(self.polygons, start=1):
            perimeter = calculate_perimeter(polygon)
            
            # Section header
            if poly_idx > 1:
                row += 2  # Add space between shapes
            
            ws.cell(row=row, column=1).value = f"SHAPE #{poly_idx}"
            ws.cell(row=row, column=1).font = Font(bold=True, size=14, color="FF0000")
            row += 1
            
            # Property headers
            headers = ["Property", "Value"]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1
            
            # Write properties
            row = self._write_shape_properties(ws, row, polygon, perimeter)
            
            # Write points
            row = self._write_shape_points(ws, row, polygon)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 15
    
    def _write_shape_properties(self, ws, row, polygon, perimeter):
        """Write shape properties to detail sheet"""
        # Shape ID (first row - most important)
        ws.cell(row=row, column=1).value = "Shape ID"
        ws.cell(row=row, column=2).value = polygon.get('id', 'N/A')
        ws.cell(row=row, column=2).font = Font(bold=True, color="FF0000", size=12)
        
        # Add "Find in CAD" button/link in column 3
        shape_id = polygon.get('id', '')
        if shape_id:
            # Use VBScript file in scripts subfolder for hidden execution (no console window)
            scripts_folder = os.path.join(self.output_folder, "scripts")
            vbs_file = os.path.join(scripts_folder, f"find_{shape_id}.vbs")
            
            cell = ws.cell(row=row, column=3)
            cell.value = "🔍 Find in CAD"
            cell.font = Font(bold=True, color="FFFFFF", size=10, underline="single")
            cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Add hyperlink to VBScript file (runs hidden)
            cell.hyperlink = vbs_file
            
            # Add comment with instructions
            from openpyxl.comments import Comment
            comment = Comment(
                f"Click to find this shape in AutoCAD\nShape ID: {shape_id}\n\nInstructions:\n1. Make sure AutoCAD is open with the drawing\n2. Click this button\n3. Shape will be highlighted in yellow",
                "CAD Inspector"
            )
            cell.comment = comment
        row += 1
        
        ws.cell(row=row, column=1).value = "Name"
        ws.cell(row=row, column=2).value = polygon.get('name', 'Unknown')
        ws.cell(row=row, column=2).font = Font(bold=True, color="0000FF")
        row += 1
        
        ws.cell(row=row, column=1).value = "Type"
        ws.cell(row=row, column=2).value = polygon.get('type', 'N/A')
        row += 1
        
        ws.cell(row=row, column=1).value = "Point Count"
        ws.cell(row=row, column=2).value = polygon.get('point_count', 0)
        row += 1
        
        ws.cell(row=row, column=1).value = "Area (sq units)"
        ws.cell(row=row, column=2).value = round(polygon.get('area', 0), 2)
        row += 1
        
        ws.cell(row=row, column=1).value = "Perimeter (units)"
        ws.cell(row=row, column=2).value = round(perimeter, 2)
        row += 1
        
        # Add radius for circles
        if polygon.get('radius'):
            ws.cell(row=row, column=1).value = "Radius"
            ws.cell(row=row, column=2).value = round(polygon.get('radius', 0), 2)
            row += 1
        
        centroid = polygon.get('centroid', [0, 0])
        ws.cell(row=row, column=1).value = "Centroid X"
        ws.cell(row=row, column=2).value = round(centroid[0], 2)
        row += 1
        
        ws.cell(row=row, column=1).value = "Centroid Y"
        ws.cell(row=row, column=2).value = round(centroid[1], 2)
        row += 1
        
        ws.cell(row=row, column=1).value = "Drawing"
        ws.cell(row=row, column=2).value = polygon.get('drawing', 'N/A')
        row += 1
        
        ws.cell(row=row, column=1).value = "Timestamp"
        ws.cell(row=row, column=2).value = polygon.get('timestamp', 'N/A')
        row += 2
        
        return row
    
    def _write_shape_points(self, ws, row, polygon):
        """Write shape points to detail sheet"""
        ws.cell(row=row, column=1).value = "Points:"
        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        
        # Points header
        point_headers = ["Point #", "X", "Y", "Z"]
        for col, header in enumerate(point_headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = Font(bold=True)
        row += 1
        
        # Points data
        points = polygon.get('points', [])
        for idx, point in enumerate(points, start=1):
            ws.cell(row=row, column=1).value = idx
            ws.cell(row=row, column=2).value = round(point[0], 2)
            ws.cell(row=row, column=3).value = round(point[1], 2)
            ws.cell(row=row, column=4).value = round(point[2], 2) if len(point) > 2 else 0
            row += 1
        
        return row
